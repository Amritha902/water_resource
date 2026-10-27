"""Build the final report from one source: Word (.docx), and Markdown for the repository.

    python tools/build_report.py --docx OUT.docx --theirs THEIRS_DIR     # private copy, with published figures
    python tools/build_report.py --md docs/REPORT.md                     # repository copy, our figures only

The PDF is made from the .docx with LibreOffice, so the two are identical.
All numbers are read from results/.  The published figures (THEIRS_DIR) are
copyrighted and are only placed in the private copy.
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from compare_pdf import FIGS  # noqa: E402
from submission_pdf import TITLES, J, p1_numbers, p2_rows, p3_numbers, pct  # noqa: E402

REPRO = ROOT / "docs" / "figures" / "reproduced"
OURS = ROOT / "docs" / "figures" / "ours"


# ---------------------------------------------------------------- content
def H(level, text): return ("h", level, text)
def P(text): return ("p", text)
def B(items): return ("ul", items)
def T(head, rows, note=None): return ("table", head, rows, note)
def PAIR(paper, n, cap, note): return ("pair", paper, n, cap, note)
def FIG(path, cap): return ("fig", path, cap)
def BOX(text): return ("box", text)
def PB(): return ("pb",)


def g3(v): return f"{v:.3g}"


def build_blocks():
    import json
    notes = json.loads((ROOT / "docs" / "comparison_notes.json").read_text())
    clean = lambda s: re.sub(r"<[^>]+>", "", s)
    p1, p3, p2 = p1_numbers(), p3_numbers(), p2_rows()
    pub2 = J("published_paper2.json")

    def p2m(w, m, k):
        v = [x[k] for x in p2 if x["weather"] == w and x["method"] == m]
        return float(np.mean(v)) if v else float("nan")

    o1 = [pct(r["us"][0], r["pid"][0]) for r in p1[:3]]
    o3 = [pct(r["us"], r["pid"]) for r in p3]
    b = []
    b.append(("title", "Reimplementing three reinforcement-learning controllers for wastewater treatment on the "
              "official BSM1 benchmark, and a retuned expert prior that improves on two of them"))
    b.append(P(f"Amritha · {date.today():%d %B %Y}"))
    b.append(H(2, "Abstract"))
    b.append(P("Three recent papers control the Benchmark Simulation Model No. 1 (BSM1) with learning controllers "
               "built on a PID prior: a multiagent adaptive critic (paper 1), a DDPG agent that lets dissolved oxygen "
               "float (paper 2) and an integrated online Q-learning scheme (paper 3). We reimplemented all three "
               "exactly as printed on the official BSM1 plant with the official dry, rain and storm influent, and "
               "redrew every figure of the three papers from our runs. The plant and test conditions are largely "
               "confirmed: the oxygen-loop PID results of papers 1 and 3 and the PID energy of paper 2 are reproduced "
               "within 1–10 % (one case 16 %), while paper 1's nitrate-loop PID error comes out about half the "
               "published value. The learning controllers are not reproduced: paper 1's controller stays equal to its "
               "PID prior (IAE within 1 %) in 80 of 80 runs, and paper 3's is unstable as printed; paper 2's DDPG reproduces the "
               "direction of its energy saving, but with more ammonium and a wide spread between seeds. We show that "
               "the PID prior used by papers 1 and 3 is weakly tuned: raising its gain, with the factor chosen on dry "
               f"weather only and a margin from the stability limit, reduces the tracking error by "
               f"{min(o1):.0f}–{max(o1):.0f} % (paper 1, oxygen) and {min(o3):.0f}–{max(o3):.0f} % (paper 3) in "
               "every weather, including conditions not used for the choice. For paper 2, our constrained agent held "
               "to the PID's own compliance does not improve on the PID; we report this negative result."))

    b.append(H(2, "1. The three base papers"))
    for i in (1, 2, 3):
        a, t, v = TITLES[i]
        b.append(BOX(f"Paper {i}. {t}\n{a}. {v}"))
    b.append(H(2, "2. What was done"))
    b.append(B(["Implemented each algorithm exactly as printed, with every value given in the paper.",
                "Ran each paper's own experiments on the official BSM1 plant (the bsm2-python port of the IWA "
                "benchmark code) with the official dry, rain and storm influent files.",
                "Redrew every figure of the three papers (24 in all) from these runs and compared every results "
                "table with ours.",
                "Tested an improved controller for each paper under identical conditions."]))
    b.append(P("Common protocol: 14 days at 45 s control periods (26 880 periods); before each run the plant runs 14 "
               "days on dry weather under the paper's PID, so every controller starts from the same state; "
               "IAE = mean|e|, ISE = mean e², DEVmax = max|e|. Repeated runs differ only in the network "
               "initialisation (and the noise sequence). The plant, the data and every setting are listed in "
               "Appendix A; the code in Appendix B."))

    # ------------------------------------------------------------ paper 1
    b.append(PB())
    b.append(H(2, "3. Paper 1"))
    a, t, v = TITLES[1]
    b.append(BOX(f"Paper 1. {t}\n{a}. {v}"))
    b.append(P("Method. Two agents, one per loop (KLa,5 → S_O,5 and Q_a → S_NO,2), each with an action network "
               "Δa = π(z) and a critic Q(z, Δa, a). The applied input is u = a + b, where b is an incremental PID "
               "(Kp = diag(100, 100 000), Ki = diag(20, 30 000), Kd = diag(10, 1000)) and a = a + Δa. Table I of the "
               "paper gives every weight, learning rate and network size; we used them unchanged."))
    b.append(P("A discrepancy in the paper. The text gives set points r = [1, 2] for [S_O,5, S_NO,2], but Figs. 4–6 "
               "and 8 show S_O,5 held at 2 and S_NO,2 at 1. With [2, 1] our PID reproduces the paper's PID oxygen "
               "results (dry IAE 1.140·10⁻² against 1.14·10⁻²); with [1, 2] it does not. We use [2, 1]. Our "
               "nitrate-loop PID error is about half the paper's in every weather; the paper's nitrate controller, "
               "sensor or actuator must differ from what the text describes, and we could not find what."))
    b.append(H(3, "3.1 Figures: published and reimplemented"))
    for n, cap in FIGS[1][1].items():
        b.append(PAIR(1, n, cap, clean(notes.get(f"p1_fig{n}", ""))))
    rows = []
    for r in p1:
        rows.append([r["lab"], "PID (paper's prior)", g3(r["ref"]["PID"]["IAE_SO5"]), g3(r["rep"][0]),
                     g3(r["ref"]["PID"]["IAE_SNO2"]), g3(r["rep"][1])])
        rows.append(["", "MAACC", g3(r["ref"]["MAACC"]["IAE_SO5"]), g3(r["re"][0]),
                     g3(r["ref"]["MAACC"]["IAE_SNO2"]), g3(r["re"][1])])
    b.append(H(3, "3.2 Tables II and III against our reimplementation"))
    b.append(T(["Case", "Method", "IAE S_O,5 paper", "IAE S_O,5 ours", "IAE S_NO,2 paper", "IAE S_NO,2 ours"], rows,
               "Ours: mean of 20 runs."))
    b.append(P("Result: not reproduced. Run as printed, MAACC stays equal to its PID prior in all 80 runs (IAE within 0.8 %). With the "
               "Table I weights in raw units the utility is about 10⁻⁴, so at lc = 0.01 and la = 0.005 neither "
               "network moves (the inputs change by under 1 %). Twelve other initial weight scales, the only setting "
               "the paper leaves open, either leave it equal to PID or make it diverge. Under the stated actuator "
               "noise our PID error roughly doubles, whereas the paper's barely changes, so that experiment cannot "
               "be matched."))
    b.append(H(3, "3.3 Our improvement: a retuned expert prior"))
    b.append(P("The learned correction in paper 1 is meant to improve on the PID prior. On the official plant the "
               "prior itself has large headroom: scaling all its gains by a factor g reduces the error until the "
               "oxygen loop enters a limit cycle at g = 5 (dry weather). We choose g = 3, the largest step below 70 % "
               "of that limit, on dry weather only, and apply it unchanged to rain, storm and the noise test."))
    b.append(FIG(OURS / "ours_p1_dry.png", "Paper 1 set-up, dry weather: the paper's PID prior and our retuned prior."))
    rows = [[r["lab"], g3(r["pid"][0]), g3(r["re"][0]), f"{g3(r['us'][0])} (−{pct(r['us'][0], r['pid'][0]):.0f} %)",
             g3(r["pid"][1]), g3(r["re"][1]), f"{g3(r['us'][1])} (−{pct(r['us'][1], r['pid'][1]):.0f} %)"] for r in p1]
    b.append(T(["Case", "S_O,5: PID", "MAACC (reimpl.)", "Ours", "S_NO,2: PID", "MAACC (reimpl.)", "Ours"], rows,
               "IAE. Percentages: reduction relative to the paper's PID prior. Noise case: mean of 20 noise sequences. "
               "The gain was chosen on dry weather, so rain, storm and noise are unseen conditions."))

    # ------------------------------------------------------------ paper 2
    b.append(PB())
    b.append(H(2, "4. Paper 2"))
    a, t, v = TITLES[2]
    b.append(BOX(f"Paper 2. {t}\n{a}. {v}"))
    b.append(P("Method. A DDPG agent changes KLa,5 directly, with no DO set point, using the state [S_S,5, S_O,5, "
               "S_NH,5] and two rewards: A penalises effluent limit violations, B adds aeration energy. The agent "
               "learns on week 1 and is evaluated on week 2 against a PID (DO = 2 mg/L) and a fuzzy controller. The "
               "published networks (3-256-256-1) and hyperparameters (γ 0.99, batch 256, τ 0.001, learning rates "
               "10⁻⁴ and 3·10⁻⁴) are used. Three corrections are applied because, as published, the agent did not "
               "learn reliably on our first plant: the actuator setting is added to the state, each action is held 15 "
               "minutes, and the reward is scaled by 1 − γ. Run exactly as published (one seed, dry) on the official "
               "plant, it shuts the air almost off during training and ends at a similar 10.7 % saving."))
    b.append(H(3, "4.1 Figures: published and reimplemented"))
    for n, cap in FIGS[2][1].items():
        b.append(PAIR(2, n, cap, clean(notes.get(f"p2_fig{n}", ""))))
    rows = []
    for w, lab in (("dry", "Dry"), ("rain", "Rain"), ("storm", "Storm")):
        pid = p2m(w, "PID", "AE")
        for m in ("PID", "Fuzzy", "DDPG-A", "DDPG-B"):
            ref = pub2[w][m]
            rows.append([lab if m == "PID" else "", m, f"{ref['AE']:.0f}", f"{p2m(w, m, 'AE'):.0f}",
                         f"{pct(ref['AE'], pub2[w]['PID']['AE']):.1f}", f"{pct(p2m(w, m, 'AE'), pid):.1f}",
                         f"{ref['S_NH']:.2f}", f"{p2m(w, m, 'S_NH'):.2f}", f"{ref['EQ']:.0f}", f"{p2m(w, m, 'EQ'):.0f}"])
    b.append(H(3, "4.2 Tables 2–4 against our reimplementation"))
    b.append(T(["Weather", "Method", "AE paper", "AE ours", "Saving % paper", "Saving % ours", "S_NH,e paper",
                "S_NH,e ours", "EQ paper", "EQ ours"], rows,
               "AE in kWh/d, S_NH,e in g N/m³, evaluation week. DDPG rows: mean of 4 runs (dry, rain) and 2 runs (storm)."))
    b.append(P("Result: reproduced in direction. PID and fuzzy match the paper within about 1 % on energy and effluent "
               "index, and DDPG-B saves energy in every weather, as in the paper. The size differs: our DDPG-B saves "
               "more and lets ammonium rise more, and the spread between seeds is large (dry: 5.6–15.7 %). The paper "
               "reports one value per case without seeds. Not reported by the paper: with DDPG-B the effluent "
               f"ammonium is above its 4 g N/m³ limit {100 * p2m('dry', 'DDPG-B', 'viol_NH'):.0f} % of the dry "
               f"week, against {100 * p2m('dry', 'PID', 'viol_NH'):.0f} % with the PID."))
    b.append(H(3, "4.3 Our attempt at an improvement: an agent held to the PID's own compliance"))
    b.append(P("Our agent (PANDA-RL: an influent forecast as input, a critic that optimises the worst 30 % of "
               "outcomes, and Lagrange multipliers that hold the share of time above the ammonium and total nitrogen "
               "limits to a budget) is given budgets close to the PID's own violation rates in dry weather (PID: "
               "ammonium 21 %, total nitrogen 18 % of the week; budgets used: 20 % and 18 %)."))
    b.append(FIG(OURS / "ours_p2_tradeoff.png", "Energy saved against the share of the evaluation week above the "
                 "4 g N/m³ ammonium limit; each point is one run."))
    rows = []
    for w, lab in (("dry", "Dry"), ("rain", "Rain"), ("storm", "Storm")):
        pid = p2m(w, "PID", "AE")
        for m, name in (("PID", "PID"), ("DDPG-B", "DDPG-B (paper 2)"), ("PANDA-RL-pidbudget", "Ours: constrained PANDA-RL")):
            rows.append([lab if m == "PID" else "", name, f"{pct(p2m(w, m, 'AE'), pid):.1f}",
                         f"{100 * p2m(w, m, 'viol_NH'):.0f}", f"{100 * p2m(w, m, 'viol_TN'):.0f}", f"{p2m(w, m, 'EQ'):.0f}"])
    b.append(T(["Weather", "Method", "Energy saved vs PID (%)", "Time NH4 > 4 (%)", "Time N_tot > 18 (%)", "EQ"], rows,
               "Means of runs (ours: 2 per weather)."))
    b.append(P("Result: no improvement. Held to the PID's compliance, the agent stays close to the PID: it saves at "
               "most 2.5 % energy (dry) and in rain and storm none, while still breaking the ammonium limit somewhat "
               "more often than the PID (26–42 % of the week against 21–34 %). DDPG-B's larger savings are bought "
               "with roughly twice the PID's violations. On the official plant we therefore have no controller that "
               "saves energy at the PID's compliance."))

    # ------------------------------------------------------------ paper 3
    b.append(PB())
    b.append(H(2, "5. Paper 3"))
    a, t, v = TITLES[3]
    b.append(BOX(f"Paper 3. {t}\n{a}. {v}"))
    b.append(P("Method. u = ω·α(e) + β(e): a PID prior α (Kp = 100, Ki = 20, Kd = 1) plus an action network β, "
               "both included in the critic Q(e, u); utility e² + 0.01u², γ = 1, 15 hidden units, lc = 0.01, "
               "la = 0.005, 50 critic and 60 action-network training steps per period. Fixed (r = 2 mg/L) and "
               "dynamic (2 / 2.2 / 1.8 mg/L) references; priors K1–K3 with ω = 0.99^k."))
    b.append(H(3, "5.1 Figures: published and reimplemented"))
    for n, cap in FIGS[3][1].items():
        b.append(PAIR(3, n, cap, clean(notes.get(f"p3_fig{n}", ""))))
    rows = [[r["lab"], g3(r["ref"]["PID"][0]), g3(r["pid"]), g3(r["ref"]["IOQL"][0]), g3(r["re_med"]),
             " ".join(f"{v:.2g}" for v in r["re_all"])] for r in p3]
    b.append(H(3, "5.2 Tables I and III against our reimplementation"))
    b.append(T(["Case", "PID paper", "PID ours", "IOQL paper", "IOQL ours (median)", "IOQL ours, each run"], rows,
               "IAE of S_O,5. Ours: 3 runs per case."))
    b.append(P("Result: PID reproduced within 5 % (dynamic storm: 16 % lower); IOQL not reproduced. As printed it "
               "either stays at the PID level or drives the aeration down until S_O,5 collapses. The utility term "
               "0.01·u² (about 100–400 for KLa,5 = 100–200 d⁻¹) outweighs e² (about 10⁻⁴) by six orders of "
               "magnitude, so minimising Q lowers the air. The prior-policy experiments (Tables II and IV, priors "
               "K1–K3 with ω = 0.99^k) fail in the same way."))
    b.append(H(3, "5.3 Our improvement: a retuned expert prior"))
    b.append(P("The same approach as for paper 1: the prior's gain is raised by a factor chosen on dry weather with "
               "a fixed reference only (limit cycle at 6, chosen factor 4) and applied unchanged to rain, storm and "
               "the dynamic reference."))
    b.append(FIG(OURS / "ours_p3_dynamic.png", "Paper 3 set-up, dynamic reference: the paper's PID prior and our "
                 "retuned prior."))
    rows = [[r["lab"], g3(r["pid"]), g3(r["re_med"]), f"{g3(r['us'])} (−{pct(r['us'], r['pid']):.0f} %)",
             f"{g3(r['ref']['IOQL'][0])} (−{pct(r['ref']['IOQL'][0], r['ref']['PID'][0]):.0f} %)"] for r in p3]
    b.append(T(["Case", "PID", "IOQL (reimpl., median)", "Ours", "IOQL as published"], rows, "IAE of S_O,5."))

    # ------------------------------------------------------------ summary
    b.append(PB())
    b.append(H(2, "6. Overall comparison"))
    b.append(FIG(OURS / "ours_summary.png", "Reduction of the tracking error relative to each paper's PID prior: as "
                 "claimed in the paper, as obtained by reimplementing the paper exactly, and with our retuned prior."))
    b.append(P("What is supported. On the official benchmark, the published learning controllers of papers 1 and 3 "
               "do not deliver their reported improvements when implemented as printed; our retuned prior delivers "
               f"{min(o1 + o3):.0f}–{max(o1 + o3):.0f} % lower oxygen error in every weather and reference, including "
               "conditions it was not tuned on, without any learning. For paper 2 the energy saving of DDPG is "
               "reproduced in direction, and its cost in effluent compliance, not reported by the paper, is measured; "
               "our constrained agent did not improve on the PID."))
    b.append(P("What is not claimed. Our gains are smaller than the reductions claimed by papers 1 and 3 (oxygen 92–95 %, nitrate 71–76 %), which we could "
               "not obtain with any setting the papers allow. The retuned prior is a fixed controller chosen from "
               f"dry-weather runs, not a learning method; under actuator noise its gain shrinks (oxygen "
               f"−{pct(p1[3]['us'][0], p1[3]['pid'][0]):.0f} %, nitrate −{pct(p1[3]['us'][1], p1[3]['pid'][1]):.0f} %). "
               "All results are from simulation of one benchmark plant."))
    b.append(P("Findings that concern the papers themselves: paper 1 states set points [1, 2] but its figures and PID "
               "results correspond to [2, 1]; paper 1's critic update (eq. 15) grows without bound for slowly varying "
               "states; paper 1's noise experiment cannot be matched; paper 3's utility makes its learner lower the "
               "air; none of the three gives initial weights, input scaling or code."))

    # ------------------------------------------------------------ appendices
    b.append(PB())
    b.append(H(2, "Appendix A. Simulation data and settings"))
    b.extend(md_blocks((ROOT / "docs" / "DATASET.md").read_text(), demote=True, prefix="A."))
    b.append(PB())
    b.append(H(2, "Appendix B. Implementation"))
    b.append(T(["Item", "File"], [
        ["Official BSM1 plant, influent loading and correction", "src/wwtp/official/plant.py"],
        ["Paper 1, MAACC exactly as printed", "src/wwtp/maacc/exact.py"],
        ["Paper 1 experiments (Tables II, III; gain sweep; retuned prior)", "experiments/20_exact_paper1.py"],
        ["Paper 2, environment, DDPG, rewards, comparators", "src/wwtp/rl/env.py, ddpg.py, rewards.py, baselines.py"],
        ["Paper 2 experiments (official plant)", "experiments/06_aeration_benchmark.py --official"],
        ["Paper 2, Fig. 4 (effluent vs DO set point)", "experiments/26_fig4_official.py"],
        ["Paper 2, constrained agent (PANDA-RL)", "src/wwtp/rl/panda_rl.py"],
        ["Paper 3, IOQL exactly as printed", "src/wwtp/ioql/ioql.py"],
        ["Paper 3 experiments (Tables I–IV; retuned prior)", "experiments/23_ioql.py"],
        ["Redrawn diagrams; reproduced result figures; our figures",
         "experiments/24_diagrams.py, 25_paper_figures.py, 27_ours_figures.py"],
        ["Published values typed in from the papers", "results/published_paper1.json, _paper2.json, _paper3.json"],
        ["This report", "tools/build_report.py"],
        ["Tests (82)", "tests/"],
    ]))
    b.append(P("Commands, from the repository root:"))
    b.append(("code", "pip install -r requirements.txt\n"
              "python -m pytest tests -q\n"
              "python experiments/20_exact_paper1.py --seeds 20 --noise --critics paper --tag exact2\n"
              "python experiments/20_exact_paper1.py --controllers PIDx3 --seeds 1 --weathers dry rain storm --tag ours_p1_clean\n"
              "python experiments/23_ioql.py --seeds 3\n"
              "python experiments/23_ioql.py --methods PID PIDx4 --tag ours_p3\n"
              "python experiments/06_aeration_benchmark.py --official --weathers dry rain storm --methods PID Fuzzy DDPG-A DDPG-B\n"
              "python experiments/26_fig4_official.py\n"
              "python experiments/24_diagrams.py && python experiments/25_paper_figures.py && python experiments/27_ours_figures.py\n"
              "python tools/build_report.py --md docs/REPORT.md"))
    b.append(H(2, "Appendix C. Runs behind this report"))
    inv = [["Paper 1 as printed (dry, rain, storm, dry + noise; 20 each, plus PID)", "103"],
           ["Paper 1 initial-weight sweep", "24"],
           ["Paper 1 gain sweep (dry) and retuned prior (all cases)", "45"],
           ["Paper 2 PID, fuzzy, DDPG-A, DDPG-B (dry, rain, storm)", "26"],
           ["Paper 2 as published (dry)", "1"],
           ["Paper 2 constrained agent (dry, rain, storm, 2 each)", "6"],
           ["Paper 2 Fig. 4 steady-state set points", "8"],
           ["Paper 3 IOQL (fixed and dynamic; PID, IOQL, K1–K3; 3 each)", "78"],
           ["Paper 3 gain sweep (dry) and retuned prior (all cases)", "18"]]
    b.append(T(["Experiment", "Runs"], inv))
    b.append(H(2, "References"))
    for i in (1, 2, 3):
        a, t, v = TITLES[i]
        b.append(P(f"[{i}] {a}, \"{t}\", {v}."))
    b.append(P("[4] J. Alex et al., Benchmark Simulation Model no. 1 (BSM1), IWA Task Group on Benchmarking of "
               "Control Strategies for WWTPs, Lund University, 2008."))
    b.append(P("[5] bsm2-python 0.0.16, FAU-EVT, Python implementation of the IWA benchmark simulation models (BSD-3)."))
    return b


def md_blocks(text: str, demote=False, prefix=""):
    """Tiny Markdown reader for docs/DATASET.md: headings, paragraphs, lists, tables."""
    out, para, lst, tab = [], [], [], []
    inline = lambda s: re.sub(r"`([^`]*)`", r"\1", s).replace("**", "")

    def flush():
        nonlocal para, lst, tab
        if para:
            out.append(P(inline(" ".join(para))))
        if lst:
            out.append(B([inline(x) for x in lst]))
        if tab:
            rows = [[inline(c.strip()) for c in r.strip().strip("|").split("|")] for r in tab]
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
            out.append(T(rows[0], rows[1:]))
        para, lst, tab = [], [], []

    for line in text.splitlines():
        s = line.rstrip()
        if s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            if level == 1:
                continue
            head = s.lstrip("#").strip()
            out.append(H(min(level + (1 if demote else 0), 4), prefix + head if head[:1].isdigit() else head))
        elif s.startswith("|"):
            if para or lst:
                p_, l_ = para, lst
                para, lst = [], []
                if p_:
                    out.append(P(inline(" ".join(p_))))
                if l_:
                    out.append(B([inline(x) for x in l_]))
            tab.append(s)
        elif re.match(r"^\s*(- |\d+\. )", s):
            if tab or para:
                flush()
            lst.append(re.sub(r"^\s*(- |\d+\. )", "", s))
        elif s.strip() == "":
            flush()
        elif lst and not para and s.startswith("  "):
            lst[-1] += " " + s.strip()
        else:
            if tab:
                flush()
            para.append(s.strip())
    flush()
    return out


# ---------------------------------------------------------------- Word
def to_docx(blocks, out: Path, theirs: Path | None):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for name, size in (("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 11.5), ("Heading 4", 10.5)):
        hs = doc.styles[name]
        hs.font.name, hs.font.size, hs.font.color.rgb = "Arial", Pt(size), RGBColor(0x1F, 0x1F, 0x1E)
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    width = sec.page_width - sec.left_margin - sec.right_margin

    def shade(cell, hexcol):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcol)
        tcPr.append(sh)

    def caption(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.italic, r.font.size = True, Pt(9)
        return p

    for blk in blocks:
        k = blk[0]
        if k == "title":
            p = doc.add_paragraph()
            r = p.add_run(blk[1])
            r.bold, r.font.size, r.font.name = True, Pt(17), "Arial"
        elif k == "h":
            doc.add_heading(blk[2], level=min(blk[1] - 1, 3)).paragraph_format.keep_with_next = True
        elif k == "p":
            p = doc.add_paragraph(blk[1])
            p.paragraph_format.space_after = Pt(5)
        elif k == "ul":
            for it in blk[1]:
                doc.add_paragraph(it, style="List Bullet")
        elif k == "box":
            t = doc.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            c = t.cell(0, 0)
            shade(c, "F1F0EC")
            first, rest = blk[1].split("\n", 1)
            c.paragraphs[0].add_run(first).bold = True
            c.add_paragraph(rest).runs[0].font.size = Pt(9.5)
            doc.add_paragraph()
        elif k == "table":
            head, rows, note = blk[1], blk[2], blk[3]
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Light Grid Accent 1" if False else "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(head):
                c = t.rows[0].cells[i]
                c.text = ""
                run = c.paragraphs[0].add_run(h)
                run.bold, run.font.size = True, Pt(8.5)
                shade(c, "E8E7E2")
            for r in rows:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = ""
                    run = cells[i].paragraphs[0].add_run(str(v))
                    run.font.size = Pt(8.5)
            if len(rows) < 30:                       # keep short tables on one page
                for row in t.rows[:-1]:
                    for c in row.cells:
                        for q in c.paragraphs:
                            q.paragraph_format.keep_with_next = True
            if note:
                caption(note)
            else:
                doc.add_paragraph()
        elif k == "pair":
            paper, n, cap, note = blk[1], blk[2], blk[3], blk[4]
            doc.add_heading(f"Fig. {n}. {cap}", level=3).paragraph_format.keep_with_next = True
            t = doc.add_table(rows=2, cols=2)
            for row in t.rows:
                trPr = row._tr.get_or_add_trPr()
                cs = OxmlElement("w:cantSplit")
                trPr.append(cs)
            for i, (lab, path) in enumerate((("Published", (theirs / f"p{paper}_fig{n}.png") if theirs else None),
                                             ("Our reimplementation", REPRO / f"p{paper}_fig{n}.png"))):
                c = t.cell(0, i)
                c.text = ""
                rr = c.paragraphs[0].add_run(lab.upper())
                rr.font.size, rr.bold = Pt(7.5), True
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c.paragraphs[0].paragraph_format.keep_with_next = True
                c2 = t.cell(1, i)
                c2.text = ""
                if path is not None and path.exists():
                    c2.paragraphs[0].add_run().add_picture(str(path), width=int(width * 0.48))
                else:
                    c2.paragraphs[0].add_run("(published figure: see the paper)").italic = True
                c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption(note)
        elif k == "fig":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(blk[1]), width=int(width * 0.85))
            p.paragraph_format.keep_with_next = True
            caption(blk[2]).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif k == "code":
            for line in blk[1].splitlines():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(line)
                r.font.name, r.font.size = "Courier New", Pt(8)
        elif k == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(out)
    print("wrote", out)


# ---------------------------------------------------------------- Markdown
def to_md(blocks, out: Path):
    rel = lambda p: Path(p).relative_to(out.parent.parent) if False else "figures/" + str(Path(p).relative_to(ROOT / "docs" / "figures"))
    L = []
    for blk in blocks:
        k = blk[0]
        if k == "title":
            L += [f"# {blk[1]}", ""]
        elif k == "h":
            L += ["#" * blk[1] + " " + blk[2], ""]
        elif k == "p":
            L += [blk[1], ""]
        elif k == "ul":
            L += [f"- {x}" for x in blk[1]] + [""]
        elif k == "box":
            first, rest = blk[1].split("\n", 1)
            L += [f"> **{first}**  ", f"> {rest}", ""]
        elif k == "table":
            L += ["| " + " | ".join(blk[1]) + " |", "|" + "---|" * len(blk[1])]
            L += ["| " + " | ".join(str(c) for c in r) + " |" for r in blk[2]]
            L += ["", f"*{blk[3]}*", ""] if blk[3] else [""]
        elif k == "pair":
            paper, n, cap, note = blk[1], blk[2], blk[3], blk[4]
            L += [f"**Fig. {n}. {cap}**", "", f"![Paper {paper}, Fig. {n}, our reimplementation]"
                  f"({rel(REPRO / f'p{paper}_fig{n}.png')})", "", note,
                  "*The published figure is in the paper (not reproduced here for copyright reasons).*", ""]
        elif k == "fig":
            L += [f"![{blk[2]}]({rel(blk[1])})", "", f"*{blk[2]}*", ""]
        elif k == "code":
            L += ["```", blk[1], "```", ""]
    out.write_text("\n".join(L))
    print("wrote", out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx")
    ap.add_argument("--theirs")
    ap.add_argument("--md")
    ap.add_argument("--dataset-docx", help="docs/DATASET.md as a Word file")
    a = ap.parse_args()
    if a.dataset_docx:
        text = (ROOT / "docs" / "DATASET.md").read_text()
        to_docx([("title", text.splitlines()[0].lstrip("# "))] + md_blocks(text), Path(a.dataset_docx), None)
    blocks = build_blocks()
    if a.docx:
        to_docx(blocks, Path(a.docx), Path(a.theirs) if a.theirs else None)
    if a.md:
        to_md(blocks, Path(a.md))
